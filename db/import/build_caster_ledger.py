# -*- coding: utf-8 -*-
"""Caster(Web Caster) 계정·퍼미션·티켓 대장 → web/data/caster-ledger.json  `PC-103`

원본: db/source/Caster_계정_퍼미션_티켓.xlsx (개발팀 관리 대장)
  · 시트 1「Web Caster 권한 설정」 = 권한 20종 정의 + B2B 기본값(소리펜 / 필기펜)
  · 시트 2~ = 고객사 1곳 = 시트 1장. 계정(email)·사용자 이름·사용기간·App key·
             티켓(코드 종류별 S/O/B/P 범위)·Permission(키 → true/false)

출력(JSON):
  perms    : [{key, group, desc, sound, write}]              — 권한 정의 20종
  accounts : [{sheet, company, companyId, id, name, since, until, appKey,
               perms:{key:bool}, ranges:[{pt,section,owner,bookStart,bookEnd,pageStart,pageEnd,count,note}],
               note}]

⚠ 대장의 **비밀번호는 저장하지 않는다** — 저장소에 평문 비밀번호를 두지 않는다.
  화면에는 「대장 참조」로 표시한다.

실행: build_all_sources.py 다음에 (회사 id 매칭에 seed-customers.json 을 쓴다)
"""
import io
import json
import os
import re
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "..", "source", "Caster_계정_퍼미션_티켓.xlsx")
WEB = os.path.join(HERE, "..", "..", "web", "data")

# 대장 시트명 → 고객사 관리의 회사명 (다를 때만 적는다)
COMPANY_ALIAS = {
    "Xiom HealthCare": "Xiom Healthcare-304",
    "Luginbühl & Cie SA": "Luginbühl & Cie SA (TruxReport)",
    "Liangshishu": "Liangshishu-111",
    "tranwisdom": "Tranwisdom",
    "Copy of NHN EDU": "에듀프레소",          # 시트 제목이 복사본 이름 그대로라 본문의 회사명을 쓴다
}
PT_OF = {"PDS2": "PDS2", "PDS3": "PDS3", "PDS4": "PDS4"}


def nz(s):
    return re.sub(r"[\s()（）\-_.,]", "", str(s or "").lower())


def cell(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_range(txt):
    """'0~238' → (0,238) · '3' → (3,3) · '' → None"""
    t = cell(txt).replace(" ", "").replace(",", "")
    if not t or t == "없음":
        return None
    m = re.match(r"^(\d+)~(\d+)$", t)
    if m:
        return int(m.group(1)), int(m.group(2))
    if re.match(r"^\d+$", t):
        return int(t), int(t)
    return None


def parse_ticket_text(pt, txt, count, note):
    """'section:3, owner:27, book:231, page:1~512' 를 줄마다 한 범위로"""
    out = []
    for line in cell(txt).split("\n"):
        line = line.strip()
        if not line or line == "없음":
            continue
        kv = {}
        for part in line.split(","):
            if ":" not in part:
                continue
            k, v = part.split(":", 1)
            kv[k.strip().lower()] = v.strip()
        # owner 값에 천단위 콤마가 있어 콤마로 먼저 쪼개면 깨진다 → 정규식으로 다시 잡는다
        m = {k: v for k, v in re.findall(r"(section|owner|book|page)\s*:\s*([\d~,\s]*)", line.lower())}
        sec = parse_range(m.get("section")); own = parse_range(m.get("owner"))
        bk = parse_range(m.get("book")); pg = parse_range(m.get("page"))
        out.append({
            "pt": pt,
            "section": sec[0] if sec else None, "sectionEnd": sec[1] if sec and sec[1] != sec[0] else None,
            "owner": own[0] if own else None, "ownerEnd": own[1] if own and own[1] != own[0] else None,
            "bookStart": bk[0] if bk else None, "bookEnd": bk[1] if bk else None,
            "pageStart": pg[0] if pg else None, "pageEnd": pg[1] if pg else None,
            "count": int(float(count)) if cell(count) not in ("", "None") else None,
            "note": note or "",
        })
    return out


def parse_ticket_table(pt, r):
    """표 형식: [PDS3 (N3C6), section, owner, book, page]"""
    sec = parse_range(r[1]); own = parse_range(r[2]); bk = parse_range(r[3]); pg = parse_range(r[4] if len(r) > 4 else "")
    if sec is None and own is None:
        return None
    return {
        "pt": pt,
        "section": sec[0] if sec else None, "sectionEnd": None,
        "owner": own[0] if own else None, "ownerEnd": None,
        "bookStart": bk[0] if bk else None, "bookEnd": bk[1] if bk else None,
        "pageStart": pg[0] if pg else None, "pageEnd": pg[1] if pg else None,
        "count": (bk[1] - bk[0] + 1) if bk else None,
        "note": "",
    }


def parse_date_range(txt):
    t = cell(txt)
    if not t:
        return "", ""
    if "무제한" in t:
        return "", "무제한"
    ds = re.findall(r"(\d{4})[.\-/](\d{1,2})[.\-/](\d{1,2})", t)
    f = lambda d: "%s-%02d-%02d" % (d[0], int(d[1]), int(d[2]))
    if len(ds) >= 2:
        return f(ds[0]), f(ds[1])
    if len(ds) == 1:
        return f(ds[0]), ""
    return "", ""


def pt_of(label):
    m = re.match(r"(PDS\d)", cell(label))
    return PT_OF.get(m.group(1)) if m else None


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws0 = wb.worksheets[0]

    # ── 권한 정의 20종
    perms = []
    group = ""
    for r in ws0.iter_rows(min_row=3, values_only=True):
        if not r[1]:
            continue
        group = cell(r[0]) or group
        perms.append({"key": cell(r[1]), "group": group, "desc": cell(r[2]),
                      "sound": cell(r[3]) == "O", "write": cell(r[4]) == "O"})

    # ── 회사 id 매칭
    seed = json.load(io.open(os.path.join(WEB, "seed-customers.json"), encoding="utf-8"))
    by_nz = {nz(c["name"]): c for c in seed["companies"]}

    accounts = []
    for ws in wb.worksheets[1:]:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(v is not None for v in r)]
        d = {"sheet": ws.title, "company": COMPANY_ALIAS.get(ws.title, ws.title), "companyId": 0,
             "id": "", "name": "", "since": "", "until": "", "appKey": "", "perms": {}, "ranges": [], "note": "",
             "addr": "", "phone": "", "homepage": ""}
        mode = None
        for r in rows:
            a = cell(r[0]); b = cell(r[1]) if len(r) > 1 else ""
            if a == "Permission" and b == "분류":
                mode = "perm"; continue
            if mode == "perm":
                if len(r) > 2 and cell(r[2]) and "/" in cell(r[2]):
                    d["perms"][cell(r[2])] = cell(r[3]).lower() == "true"; continue
                mode = None
            # 계정
            if a in ("계정", "계정 이메일"):
                d["id"] = (b if "@" in b else (cell(r[2]) if len(r) > 2 and "@" in cell(r[2]) else b)).strip()
                if a == "계정" and b and "@" not in b:
                    d["name"] = d["name"] or b       # 'study123' 처럼 계정명이 따로 있는 시트
                continue
            if a.startswith("비밀번호") or a.startswith("계정 비밀번호"):
                continue                               # 저장하지 않는다
            if a in ("사용자 이름", "사용자이름"):
                d["name"] = b; continue
            if a == "사용기간":
                d["since"], d["until"] = parse_date_range(b); continue
            if a == "App key":
                d["appKey"] = b; continue
            if a == "주소":
                d["addr"] = b; continue
            if a == "전화번호":
                d["phone"] = b; continue
            if a == "홈페이지 주소":
                d["homepage"] = b; continue
            # 티켓 — 서술형 (티켓 | PDS2 (G3C6) | section:… | 수량 | 비고) 또는 표형 (PDS3 (N3C6) | 3 | 450 | 0~238 | 0~511)
            label = b if a in ("", "티켓") else a
            pt = pt_of(label)
            if pt:
                if a in ("", "티켓"):
                    txt = cell(r[2]) if len(r) > 2 else ""
                    if a == "" and re.match(r"^\d+(~\d+)?$", txt):   # 표형인데 라벨이 B열에 있는 시트 (Liangshishu)
                        t = parse_ticket_table(pt, r[1:])
                        if t: d["ranges"].append(t)
                    elif "section" in txt.lower() or txt in ("", "없음"):
                        d["ranges"] += parse_ticket_text(pt, txt, r[3] if len(r) > 3 else None, cell(r[4]) if len(r) > 4 else "")
                    else:                              # 'PDS3 (N3C6)\t3\t942\t73\t0~512' 처럼 탭으로 붙은 행
                        parts = txt.split("\t")
                        if len(parts) >= 4:
                            t = parse_ticket_table(pt, [label] + parts[1:])
                            if t: d["ranges"].append(t)
                else:
                    t = parse_ticket_table(pt, r)
                    if t: d["ranges"].append(t)
                continue
            # 안내문(영업용 테스트 계정 등)
            if a and len(a) > 40 and not b:
                d["note"] = a
        co = by_nz.get(nz(d["company"])) or by_nz.get(nz(ws.title))
        if co:
            d["companyId"] = co["id"]; d["company"] = co["name"]
        if not d["id"]:
            d["id"] = ""                                 # 대장에 계정이 비어 있음(영신사)
        accounts.append(d)

    # 고객사 시트에만 나오는 권한(현대차 커스텀 익스포트)도 정의에 넣는다 — 기본값 없음(둘 다 X)
    known = {p["key"] for p in perms}
    for a in accounts:
        for k in a["perms"]:
            if k not in known:
                known.add(k)
                perms.append({"key": k, "group": k.split("/")[0], "desc": "고객사 전용 커스텀 (대장의 고객사 시트에만 있음)",
                              "sound": False, "write": False})

    out = {"perms": perms, "accounts": accounts,
           "source": "db/source/Caster_계정_퍼미션_티켓.xlsx", "note": "비밀번호는 저장하지 않는다 (PC-103)"}
    json.dump(out, open(os.path.join(WEB, "caster-ledger.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)

    print("권한 정의 %d종" % len(perms))
    for a in accounts:
        print("  %-22s co=%-4s %-32s 범위 %d · perm %d/%d · %s~%s %s" % (
            a["sheet"], a["companyId"] or "-", a["id"] or "(계정 없음)", len(a["ranges"]),
            sum(1 for v in a["perms"].values() if v), len(a["perms"]), a["since"], a["until"],
            "AppKey" if a["appKey"] else ""))
    miss = [a["sheet"] for a in accounts if not a["companyId"]]
    print("회사 미매칭:", miss)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
